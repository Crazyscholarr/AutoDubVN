"""Lịch nội dung, workbook Excel/JSON và mẫu phân tích offline."""
from __future__ import annotations

import json
import os
import re
import tempfile
import time
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence

from .analyze import ANALYSIS_PROMPT, DESCRIPTION_PROMPT, TITLE_PROMPT, analyze_record
from .records import DEFAULT_OUTPUT_DIR, normalize_record


def _thumb_text(item: Any) -> str:
    if isinstance(item, dict):
        return "Hình: %s | Màu: %s | Chữ: %s" % (
            item.get("description", ""), item.get("colors", ""), item.get("text", ""))
    return str(item or "")


CONTENT_CALENDAR_HEADERS = [
    "STT", "Ngày tạo", "Thứ", "Ca đăng", "Nhóm chủ đề",
    "Tiêu đề đề xuất", "Thời lượng mục tiêu", "Nguồn kịch bản", "Trạng thái",
]


def _calendar_topic(record: Dict[str, Any]) -> str:
    themes = record.get("themes") or []
    if isinstance(themes, str):
        themes = [x.strip() for x in re.split(r"[,;\n]+", themes) if x.strip()]
    return str(record.get("primary_genre") or (themes[0] if themes else "") or
               record.get("series") or "").strip()


def _calendar_title(record: Dict[str, Any], final_title: str = "") -> str:
    titles = list(record.get("titles") or [])
    return str(final_title or record.get("last_used_title") or
               (titles[0] if titles else "") or record.get("title_localized") or
               record.get("title_original") or "").strip()


def _calendar_status(record: Dict[str, Any], default: str = "Chưa làm") -> str:
    raw = str(record.get("status") or "").strip()
    mapping = {
        "new": "Chưa làm", "downloaded": "Chưa làm", "analyzed": "Chưa làm",
        "pending": "Chưa làm", "used": "Đang viết", "writing": "Đang viết",
        "rendering": "Đang dựng", "rendered": "Đã xuất video",
        "published": "Đã đăng", "skipped": "Tạm hoãn",
    }
    return mapping.get(raw.casefold(), raw or default)


def _style_content_calendar(sheet, max_row: int) -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    navy, cyan = "0B1626", "14D8D4"
    header_fill = PatternFill("solid", fgColor=navy)
    thin = Side(style="thin", color="D7E1EA")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=cyan))
    sheet.row_dimensions[1].height = 34
    widths = (8, 14, 15, 25, 27, 68, 22, 26, 22, 22, 48)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=max(2, max_row), min_col=1, max_col=9):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 4)
            cell.border = Border(bottom=thin)
        row[1].number_format = "dd/mm/yyyy"
        for col in (4, 5, 6, 7, 8, 9):
            row[col - 1].fill = PatternFill("solid", fgColor="FFF7CC")
        sheet.row_dimensions[row[0].row].height = 38
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(1, max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["J"].hidden = True
    sheet.column_dimensions["K"].hidden = True
    if max_row >= 2:
        if not sheet.data_validations.dataValidation:
            dv = DataValidation(
                type="list",
                formula1='"Chưa làm,Đang viết,Đang dựng,Đã xuất video,Đã đăng,Tạm hoãn"')
            sheet.add_data_validation(dv)
            dv.add(f"I2:I{max_row}")
        else:
            sheet.data_validations.dataValidation[0].sqref = f"I2:I{max_row}"
        if not len(sheet.conditional_formatting):
            sheet.conditional_formatting.add(
                f"I2:I{max_row}", FormulaRule(formula=['$I2="Đã đăng"'],
                                                fill=PatternFill("solid", fgColor="DCFCE7")))
            sheet.conditional_formatting.add(
                f"I2:I{max_row}", FormulaRule(formula=['$I2="Tạm hoãn"'],
                                                fill=PatternFill("solid", fgColor="FEE2E2")))
    tables = list(sheet.tables.values())
    if tables:
        tables[0].ref = f"A1:I{max(2, max_row)}"
    elif max_row >= 2:
        table = Table(displayName="ContentCalendar", ref=f"A1:I{max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def append_content_calendar(record: Dict[str, Any], output_path: str,
                            final_title: str = "", video_path: str = "",
                            target_duration: str = "1h00 - 1h30",
                            status: str = "Đã xuất video",
                            created_at: Optional[datetime] = None) -> str:
    """Thêm/cập nhật một dòng lịch nội dung; cùng ID/video không bị nhân đôi."""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để xuất lịch nội dung.") from exc

    target = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    if os.path.isfile(target) and os.path.getsize(target) > 0 and zipfile.is_zipfile(target):
        wb = load_workbook(target)
    else:
        # Một lần save bị ngắt hoặc Excel/antivirus khóa file có thể để lại file
        # 0 byte/không còn là ZIP. Giữ bản hỏng để chẩn đoán rồi tự tạo lịch mới.
        if os.path.isfile(target):
            stem, ext = os.path.splitext(target)
            backup = f"{stem}.corrupt_{time.strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
            os.replace(target, backup)
        wb = Workbook()
    if "Lịch nội dung" in wb.sheetnames:
        sheet = wb["Lịch nội dung"]
    else:
        sheet = wb.active
        sheet.title = "Lịch nội dung"
        sheet.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])
    if sheet.max_row < 1 or sheet.cell(1, 1).value != "STT":
        sheet.delete_rows(1, max(1, sheet.max_row))
        sheet.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])

    record_id = str(record.get("id") or "").strip()
    absolute_video = os.path.abspath(video_path) if video_path else ""
    row_index = 0
    for index in range(2, sheet.max_row + 1):
        if ((record_id and str(sheet.cell(index, 10).value or "") == record_id) or
                (absolute_video and str(sheet.cell(index, 11).value or "") == absolute_video)):
            row_index = index
            break
    if not row_index:
        row_index = max(2, sheet.max_row + 1)
    moment = created_at or datetime.now()
    values = [
        f"=ROW()-1", moment.date(),
        (f'=CHOOSE(WEEKDAY(B{row_index},2),"Thứ Hai","Thứ Ba","Thứ Tư",'
         '"Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ Nhật")'),
        str(record.get("best_publish_time") or "20:00"), _calendar_topic(record),
        _calendar_title(record, final_title), str(target_duration or "1h00 - 1h30"),
        str(record.get("source") or "Nhập trực tiếp"),
        _calendar_status({"status": status}, default="Đã xuất video"),
        record_id, absolute_video,
    ]
    for col, value in enumerate(values, 1):
        sheet.cell(row_index, col, value)
    _style_content_calendar(sheet, sheet.max_row)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    # openpyxl ghi trực tiếp sẽ truncate file đích trước khi hoàn tất. Ghi sang
    # file tạm cùng thư mục rồi replace để một lần dừng app không phá workbook.
    handle, temporary = tempfile.mkstemp(
        prefix=".content-calendar-", suffix=".xlsx", dir=os.path.dirname(target))
    os.close(handle)
    try:
        wb.save(temporary)
        try:
            os.replace(temporary, target)
        except PermissionError:
            stem, ext = os.path.splitext(target)
            target = f"{stem}_{time.strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
            os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return target


def export_excel(records: Sequence[Dict[str, Any]], output_path: str) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu openpyxl. Hãy chạy CAI_DAT.bat hoặc pip install openpyxl.") from exc

    rows = [dict(record) for record in records]
    if not rows:
        raise ValueError("Không có ý tưởng nào để xuất Excel.")
    target = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    wb = Workbook()
    plan = wb.active
    plan.title = "Kế hoạch sản xuất"
    dashboard = wb.create_sheet("Dashboard", 0)
    calendar = wb.create_sheet("Lịch nội dung", 1)
    series_sheet = wb.create_sheet("Cụm Series")
    raw = wb.create_sheet("Dữ liệu gốc")
    config = wb.create_sheet("Cấu hình")
    prompts = wb.create_sheet("Prompt mẫu")

    navy, cyan, teal = "0B1626", "14D8D4", "0F766E"
    white, muted, pale = "FFFFFF", "60758A", "E8F7F6"
    green, amber, red = "DCFCE7", "FEF3C7", "FEE2E2"
    thin = Side(style="thin", color="CBD5E1")
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(color=white, bold=True)

    calendar.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])
    for row_index, record in enumerate(rows, 2):
        titles = list(record.get("titles") or [])
        calendar.append([
            f"=ROW()-1", datetime.now().date(),
            (f'=CHOOSE(WEEKDAY(B{row_index},2),"Thứ Hai","Thứ Ba","Thứ Tư",'
             '"Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ Nhật")'),
            record.get("best_publish_time") or "20:00", _calendar_topic(record),
            _calendar_title(record), "1h00 - 1h30", record.get("source", ""),
            _calendar_status(record), record.get("id", ""), "",
        ])
    _style_content_calendar(calendar, calendar.max_row)

    config_rows = [
        ["Tham số", "Giá trị", "Ý nghĩa"],
        ["Trọng số Hook", .45, "Tỷ trọng trong điểm ưu tiên"],
        ["Trọng số Plot Twist", .40, "Tỷ trọng trong điểm ưu tiên"],
        ["Trọng số Dễ sản xuất", .15, "Cao=10, Trung=6, Thấp=3"],
        ["Ngưỡng Nên làm", 7.2, "Điểm ưu tiên tối thiểu"],
        ["Ngưỡng Chờ đợi", 5.5, "Dưới mức này là Không nên làm"],
        ["Tốc độ đọc (từ/phút)", 150, "Dùng ước lượng thời lượng"],
    ]
    for row in config_rows:
        config.append(row)
    config.freeze_panes = "A2"
    config.column_dimensions["A"].width = 28
    config.column_dimensions["B"].width = 14
    config.column_dimensions["C"].width = 42
    for cell in config[1]:
        cell.fill, cell.font = header_fill, header_font
    for cell in config[2][1:2] + config[3][1:2] + config[4][1:2]:
        cell.number_format = "0%"
        cell.fill = PatternFill("solid", fgColor="FFF7CC")
    for row in config.iter_rows(min_row=2, max_row=7, min_col=2, max_col=2):
        row[0].fill = PatternFill("solid", fgColor="FFF7CC")

    columns = [
        "ID", "Chọn sản xuất", "Tiêu đề gốc", "Nguồn", "URL nguồn", "Tác giả",
        "Ngày đăng", "Thể loại (chính)", "Cảm xúc chủ đạo", "Điểm Hook",
        "Điểm Plot Twist", "Điểm ưu tiên", "Chủ đề", "Series",
        "Từ khóa gợi ý", "Tag gợi ý", "Số từ", "Ước lượng thời lượng (phút)",
        "Dễ dàng sản xuất", "Khuyến nghị", "Thời điểm đăng tối ưu",
    ] + [f"Tiêu đề {i}" for i in range(1, 9)] \
      + [f"Mô tả {i}" for i in range(1, 4)] \
      + [f"Thumbnail {i}" for i in range(1, 4)] \
      + ["Outline (dàn ý 3 phần)", "Các nhân vật chính", "Trạng thái sản xuất",
         "Ghi chú", "Provider AI", "Lỗi AI"]
    last_col = get_column_letter(len(columns))
    plan.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    plan["A1"] = "GỐC MÍT KỂ CHUYỆN · KẾ HOẠCH SẢN XUẤT YOUTUBE"
    plan["A1"].fill = PatternFill("solid", fgColor=navy)
    plan["A1"].font = Font(color=white, bold=True, size=16)
    plan["A1"].alignment = Alignment(horizontal="left", vertical="center")
    plan.row_dimensions[1].height = 32
    plan.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    plan["A2"] = ("Ô vàng là trường có thể chỉnh. Điểm ưu tiên và Khuyến nghị là công thức; "
                   "lọc 'Nên làm' để đưa thẳng sang AutoDubVN.")
    plan["A2"].font = Font(color=muted, italic=True)
    plan["A2"].alignment = Alignment(wrap_text=True)
    for col, label in enumerate(columns, 1):
        cell = plan.cell(4, col, label)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=cyan))
    plan.row_dimensions[4].height = 42

    hook_col, twist_col, priority_col = 10, 11, 12
    ease_col, reco_col = 19, 20
    for idx, record in enumerate(rows, 5):
        titles = list(record.get("titles") or [])
        descriptions = list(record.get("descriptions") or [])
        thumbnails = list(record.get("thumbnails") or [])
        values = [
            record.get("id", ""), "Có" if record.get("selected") else "Không",
            record.get("title_original", ""), record.get("source", ""),
            record.get("source_url", ""), record.get("author", ""),
            record.get("published_at", ""), record.get("primary_genre", ""),
            record.get("emotion", ""), int(record.get("hook_score") or 0),
            int(record.get("plot_twist_score") or 0), None,
            ", ".join(record.get("themes") or []), record.get("series", ""),
            ", ".join(record.get("keywords") or []), ", ".join(record.get("tags") or []),
            int(record.get("word_count") or 0), None,
            record.get("production_ease", "Trung"), None,
            record.get("best_publish_time", ""),
        ]
        values += (titles + [""] * 8)[:8]
        values += (descriptions + [""] * 3)[:3]
        values += ([_thumb_text(x) for x in thumbnails] + [""] * 3)[:3]
        values += [
            record.get("outline", ""), ", ".join(record.get("main_characters") or []),
            record.get("status", "Chưa làm"), record.get("notes", ""),
            record.get("analysis_provider", ""), record.get("analysis_error", ""),
        ]
        for col, value in enumerate(values, 1):
            plan.cell(idx, col, value)
        plan.cell(idx, priority_col,
                  f'=ROUND(J{idx}*\'Cấu hình\'!$B$2+K{idx}*\'Cấu hình\'!$B$3+'
                  f'IF(S{idx}="Cao",10,IF(S{idx}="Trung",6,3))*\'Cấu hình\'!$B$4,1)')
        plan.cell(idx, 18, f'=ROUND(Q{idx}/\'Cấu hình\'!$B$7,1)')
        plan.cell(idx, reco_col,
                  f'=IF(L{idx}>=\'Cấu hình\'!$B$5,"Nên làm",'
                  f'IF(L{idx}>=\'Cấu hình\'!$B$6,"Chờ đợi","Không nên làm"))')
        for col in (2, 8, 9, 10, 11, 13, 14, 15, 16, 19, 21, 38, 39):
            plan.cell(idx, col).fill = PatternFill("solid", fgColor="FFF7CC")
        for col in range(1, len(columns) + 1):
            plan.cell(idx, col).alignment = Alignment(
                vertical="top", wrap_text=col >= 3)
        plan.row_dimensions[idx].height = 48

    plan.freeze_panes = "C5"
    plan.auto_filter.ref = f"A4:{last_col}{len(rows) + 4}"
    table = Table(displayName="ProductionPlan", ref=f"A4:{last_col}{len(rows) + 4}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    plan.add_table(table)
    widths = {
        1: 20, 2: 14, 3: 34, 4: 16, 5: 32, 6: 18, 7: 16, 8: 23, 9: 18,
        10: 11, 11: 13, 12: 12, 13: 28, 14: 30, 15: 30, 16: 30, 17: 11,
        18: 17, 19: 17, 20: 15, 21: 27,
    }
    for col in range(22, 30):
        widths[col] = 42
    for col in range(30, 33):
        widths[col] = 52
    for col in range(33, 36):
        widths[col] = 50
    widths.update({36: 58, 37: 32, 38: 20, 39: 30, 40: 24, 41: 42})
    for col, width in widths.items():
        plan.column_dimensions[get_column_letter(col)].width = width
    max_row = len(rows) + 4
    dv_yes = DataValidation(type="list", formula1='"Có,Không"', allow_blank=False)
    dv_ease = DataValidation(type="list", formula1='"Cao,Trung,Thấp"')
    dv_status = DataValidation(
        type="list", formula1='"Chưa làm,Đang viết,Đã duyệt,Đang dựng,Đã đăng,Tạm hoãn"')
    plan.add_data_validation(dv_yes); dv_yes.add(f"B5:B{max_row}")
    plan.add_data_validation(dv_ease); dv_ease.add(f"S5:S{max_row}")
    plan.add_data_validation(dv_status); dv_status.add(f"AL5:AL{max_row}")
    for rng in (f"J5:J{max_row}", f"K5:K{max_row}", f"L5:L{max_row}"):
        plan.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=1, start_color="FEE2E2",
            mid_type="num", mid_value=6, mid_color="FEF3C7",
            end_type="num", end_value=10, end_color="DCFCE7"))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Nên làm\""],
                                       fill=PatternFill("solid", fgColor=green)))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Chờ đợi\""],
                                       fill=PatternFill("solid", fgColor=amber)))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Không nên làm\""],
                                       fill=PatternFill("solid", fgColor=red)))
    plan["L4"].comment = Comment(
        "Hook × trọng số + Plot Twist × trọng số + mức dễ sản xuất × trọng số.",
        "AutoDubVN")

    # Dữ liệu gốc giữ nguyên để kiểm toán nguồn.
    raw_headers = ["ID", "Tiêu đề gốc", "Nguồn", "URL nguồn", "Ngôn ngữ",
                   "Số từ", "Nội dung đang dùng", "Nội dung thô"]
    raw.append(raw_headers)
    for record in rows:
        raw.append([
            record.get("id", ""), record.get("title_original", ""),
            record.get("source", ""), record.get("source_url", ""),
            record.get("language", ""), int(record.get("word_count") or 0),
            str(record.get("content") or "")[:32700],
            str(record.get("raw_content") or "")[:32700],
        ])
    for cell in raw[1]:
        cell.fill, cell.font = header_fill, header_font
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:H{len(rows) + 1}"
    for col, width in enumerate((20, 38, 16, 35, 12, 12, 85, 85), 1):
        raw.column_dimensions[get_column_letter(col)].width = width
    for row in raw.iter_rows(min_row=2):
        row[6].alignment = row[7].alignment = Alignment(wrap_text=True, vertical="top")
        raw.row_dimensions[row[0].row].height = 72

    series_names = sorted({str(r.get("series") or "Chưa phân cụm") for r in rows})
    series_sheet.append(["Series", "Số ý tưởng", "Hook trung bình", "Twist trung bình",
                         "Chủ đề đại diện"])
    for i, name in enumerate(series_names, 2):
        series_sheet.cell(i, 1, name)
        series_sheet.cell(i, 2, f'=COUNTIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i})')
        series_sheet.cell(i, 3, f'=IFERROR(AVERAGEIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i},\'Kế hoạch sản xuất\'!$J$5:$J${max_row}),0)')
        series_sheet.cell(i, 4, f'=IFERROR(AVERAGEIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i},\'Kế hoạch sản xuất\'!$K$5:$K${max_row}),0)')
        example = next((r for r in rows if str(r.get("series") or "Chưa phân cụm") == name), {})
        series_sheet.cell(i, 5, ", ".join(example.get("themes") or []))
    for cell in series_sheet[1]:
        cell.fill, cell.font = header_fill, header_font
    series_sheet.freeze_panes = "A2"
    for col, width in enumerate((38, 15, 18, 18, 42), 1):
        series_sheet.column_dimensions[get_column_letter(col)].width = width

    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:H2")
    dashboard["A1"] = "BẢNG ƯU TIÊN SẢN XUẤT NỘI DUNG"
    dashboard["A1"].fill = PatternFill("solid", fgColor=navy)
    dashboard["A1"].font = Font(color=white, bold=True, size=18)
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")
    cards = [
        ("A4", "Tổng ý tưởng", f'=COUNTA(\'Kế hoạch sản xuất\'!$A$5:$A${max_row})'),
        ("C4", "Nên làm", f'=COUNTIF(\'Kế hoạch sản xuất\'!$T$5:$T${max_row},"Nên làm")'),
        ("E4", "Hook TB", f'=ROUND(AVERAGE(\'Kế hoạch sản xuất\'!$J$5:$J${max_row}),1)'),
        ("G4", "Twist TB", f'=ROUND(AVERAGE(\'Kế hoạch sản xuất\'!$K$5:$K${max_row}),1)'),
    ]
    for anchor, label, formula in cards:
        col = dashboard[anchor].column
        dashboard.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        dashboard.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 1)
        dashboard.cell(4, col, label)
        dashboard.cell(5, col, formula)
        dashboard.cell(4, col).fill = PatternFill("solid", fgColor=teal)
        dashboard.cell(4, col).font = Font(color=white, bold=True)
        dashboard.cell(5, col).fill = PatternFill("solid", fgColor=pale)
        dashboard.cell(5, col).font = Font(color=navy, bold=True, size=20)
        dashboard.cell(5, col).alignment = Alignment(horizontal="center", vertical="center")
    dashboard["A8"] = "Cách dùng"
    dashboard["A8"].font = Font(bold=True, size=13, color=navy)
    dashboard.merge_cells("A9:H11")
    dashboard["A9"] = (
        "1) Lọc Khuyến nghị = Nên làm. 2) Chọn một Tiêu đề AI và chỉnh ô vàng. "
        "3) Đổi Chọn sản xuất = Có. 4) Trong AutoDubVN bấm Dùng trong AI Story "
        "để chuyển tiêu đề, outline, tag và mô tả sang quy trình dựng video.")
    dashboard["A9"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard["A9"].fill = PatternFill("solid", fgColor="F8FAFC")
    dashboard["A9"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, 9):
        dashboard.column_dimensions[get_column_letter(col)].width = 16
    if series_names:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Số ý tưởng theo Series"
        chart.y_axis.title = "Series"
        chart.x_axis.title = "Số ý tưởng"
        data = Reference(series_sheet, min_col=2, min_row=1,
                         max_row=len(series_names) + 1)
        cats = Reference(series_sheet, min_col=1, min_row=2,
                         max_row=len(series_names) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 8, 17
        dashboard.add_chart(chart, "A13")

    prompts.sheet_view.showGridLines = False
    prompts.append(["Nhiệm vụ", "Prompt mẫu"])
    prompts.append(["Đánh giá + kế hoạch đa tầng", ANALYSIS_PROMPT])
    prompts.append(["Tạo 8 tiêu đề", TITLE_PROMPT])
    prompts.append(["Tạo 3 mô tả", DESCRIPTION_PROMPT])
    for cell in prompts[1]:
        cell.fill, cell.font = header_fill, header_font
    prompts.column_dimensions["A"].width = 32
    prompts.column_dimensions["B"].width = 110
    for row in prompts.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        prompts.row_dimensions[row[0].row].height = 150

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(target)
    return target


def export_plan(records: Sequence[Dict[str, Any]], output_dir: str = "",
                stem: str = "ke_hoach_noi_dung") -> Dict[str, str]:
    folder = os.path.abspath(output_dir or DEFAULT_OUTPUT_DIR)
    os.makedirs(folder, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^\w\-]+", "_", stem, flags=re.UNICODE).strip("_")
    xlsx = export_excel(records, os.path.join(folder, f"{safe_stem}_{stamp}.xlsx"))
    json_path = os.path.join(folder, f"{safe_stem}_{stamp}.json")
    with open(json_path, "w", encoding="utf-8") as handle:
        json.dump(list(records), handle, ensure_ascii=False, indent=2)
    return {"xlsx": xlsx, "json": json_path, "output_dir": folder}


def sample_records() -> List[Dict[str, Any]]:
    samples = [
        {"titleOriginal": "Ba người con tranh căn nhà của mẹ già",
         "source": "VnExpress Tâm sự", "sourceUrl": "https://vnexpress.net/tam-su",
         "language": "vi", "rawContent": (
             "Bà Hòa đã ngoài bảy mươi, sống một mình trong căn nhà cấp bốn. "
             "Ba người con chỉ trở về khi nghe tin bà định lập di chúc. Trong bữa cơm giỗ, "
             "một tờ giấy cũ bất ngờ xuất hiện khiến mọi người tranh cãi dữ dội. ") * 120},
        {"titleOriginal": "婆媳矛盾：老人被赶出家门",
         "source": "Zhihu 盐选", "sourceUrl": "https://www.zhihu.com/",
         "language": "zh", "rawContent": (
             "婆婆和儿媳因为赡养和房产发生矛盾，老人被赶出家门。后来一份遗产文件揭开真相。") * 280},
        {"titleOriginal": "AITA for refusing to pay my father's nursing home?",
         "source": "Reddit AITA", "sourceUrl": "https://www.reddit.com/r/AmItheAsshole",
         "language": "en", "rawContent": (
             "My siblings argued about who should care for our elderly father after our mother died. "
             "Everyone wanted the inheritance, but nobody agreed to visit him at the nursing home. ") * 170},
    ]
    return [analyze_record(normalize_record(row, i), use_ai=False)
            for i, row in enumerate(samples)]
